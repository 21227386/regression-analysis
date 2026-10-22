from io import BytesIO
import unittest
import numpy as np
import pandas as pd
from excel_io import read_excel, workbook_bytes
from model import example, fit, validate

class ExcelTests(unittest.TestCase):
    def test_excel_roundtrip_and_fit(self):
        df = example().drop(columns='함명')
        loaded = read_excel(BytesIO(workbook_bytes(df)))
        result = fit(loaded)
        self.assertEqual(result.reference, '지원함')
        self.assertEqual(len(result.results), 60)
        self.assertEqual(result.coefficients.set_index('함종').loc['지원함', 'D_T'], 1)
        np.testing.assert_allclose(result.results['예측중량'], fit(example()).results['예측중량'])

    def test_unknown_type_and_missing_reference(self):
        df = example()
        df.loc[0, '함종'] = '기타'
        with self.assertRaisesRegex(ValueError, '허용되지 않은 함종'): validate(df)
        df = example().query("함종 != '지원함'")
        with self.assertRaisesRegex(ValueError, '지원함'): fit(df)

    def test_select_sheet_and_blank_rows(self):
        buf = BytesIO()
        df = example()
        with pd.ExcelWriter(buf, engine='openpyxl') as writer:
            pd.DataFrame({'설명': ['자료는 두 번째 시트']}).to_excel(writer, index=False, sheet_name='설명')
            pd.concat([df.iloc[:2], pd.DataFrame([{}]), df.iloc[2:]]).to_excel(writer, index=False, sheet_name='자료')
        self.assertEqual(len(read_excel(BytesIO(buf.getvalue()), '자료')), len(df))

class SpeedLabelTests(unittest.TestCase):
    def test_not_template_header_and_roundtrip(self):
        from openpyxl import load_workbook
        content = workbook_bytes(example())
        book = load_workbook(BytesIO(content), read_only=True)
        self.assertIn('순항속력 (kn)', next(book.active.values))
        loaded = read_excel(BytesIO(content))
        self.assertIn('순항속력', loaded.columns)
        np.testing.assert_allclose(loaded['순항속력'], example()['순항속력'])

    def test_old_speed_headers_and_landing_ship_dropdown(self):
        from openpyxl import load_workbook
        frame = example()
        for header in ['순항속력', '순항속력 (kn)', '순항속력(kn)', '순항속력 (not)', '순항속력(knot)']:
            buffer = BytesIO()
            frame.rename(columns={'순항속력': header}).to_excel(buffer, index=False)
            loaded = read_excel(BytesIO(buffer.getvalue()))
            np.testing.assert_allclose(loaded['순항속력'], frame['순항속력'])
        book = load_workbook(BytesIO(workbook_bytes(frame)))
        self.assertIn('상륙함', book.active.data_validations.dataValidation[0].formula1)
